#!/usr/bin/env python3
"""Portable write-only Excel export of the frozen typed source-data inputs.

Unlike the original JavaScript authoring path, this exporter requires only
openpyxl. It preserves source values, including literal strings that begin with
Excel formula prefixes. The visual style and ZIP bytes can differ from the
frozen workbooks. Optional verification reads OOXML independently of openpyxl.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import posixpath
import re
import tempfile
import warnings
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

HERE = Path(__file__).resolve().parent
MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
DOC_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PACKAGE_REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
TAG = '{' + MAIN + '}'
HEADER_FILL = PatternFill('solid', fgColor='23485F')
HEADER_FONT = Font(name='Arial', size=10, bold=True, color='FFFFFF')
TEXT_FONT = Font(name='Arial', size=10, color='24313B')


def typed_cell(sheet, value, *, header=False):
    """Retain JSON types; strings are always literal text, never formulas."""
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError('An Excel string exceeds the 32,767-character limit')
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError('Non-finite JSON numbers must use explicit source flags')
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = 's'
    cell.font = HEADER_FONT if header else TEXT_FONT
    cell.alignment = Alignment(vertical='top', wrap_text=header)
    if header:
        cell.fill = HEADER_FILL
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.number_format = '0.###############'
    return cell


def append_typed(sheet, values, *, header=False):
    sheet.append([typed_cell(sheet, value, header=header) for value in values])


def add_readme(workbook, book):
    sheet = workbook.create_sheet('Readme')
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions['A'].width = 28
    sheet.column_dimensions['B'].width = 112
    sheet.freeze_panes = 'A5'
    append_typed(sheet, [])
    append_typed(sheet, [book['title']])
    append_typed(sheet, [])
    append_typed(sheet, [])
    for label, description in book['notes']:
        append_typed(sheet, [label, description])


def add_data(workbook, book_id, index, source):
    headers = source['headers']
    if any(not isinstance(value, str) for value in headers):
        raise ValueError('Native Excel table headers must be strings')
    if len({value.casefold() for value in headers}) != len(headers):
        raise ValueError('Native Excel table headers collide case-insensitively')
    sheet = workbook.create_sheet(source['name'])
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A6'
    for index_col in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index_col)].width = 22
    append_typed(sheet, [])
    append_typed(sheet, [source['title']])
    append_typed(sheet, [source['source']])
    append_typed(sheet, [])
    append_typed(sheet, headers, header=True)
    for row in source['rows']:
        if len(row) != len(headers):
            raise ValueError('Typed source row has an inconsistent column count')
        append_typed(sheet, row)
    reference = f'A5:{get_column_letter(len(headers))}{len(source["rows"]) + 5}'
    table_name = re.sub(r'[^A-Za-z0-9_]', '_', f'Data_{book_id}_{index}')
    table = Table(
        displayName=table_name,
        ref=reference,
        autoFilter=AutoFilter(ref=reference),
        tableColumns=[TableColumn(id=i, name=name)
                      for i, name in enumerate(headers, 1)],
        tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',
                                      showRowStripes=True),
    )
    # All table columns are supplied explicitly, as write-only mode requires.
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore',
                                message='In write-only mode you must add table columns manually')
        sheet.add_table(table)


def build_book(book, output_dir):
    destination = output_dir / book['file']
    if destination.exists():
        raise FileExistsError(f'Refusing to overwrite an existing workbook: {destination}')
    workbook = Workbook(write_only=True)
    workbook.properties.title = book['title']
    workbook.properties.subject = 'Numerical source data for reproducible social-learning simulations'
    workbook.properties.creator = 'Decision-generated credibility replication package'
    add_readme(workbook, book)
    for index, source in enumerate(book['sheets'], 1):
        add_data(workbook, book['id'], index, source)
    descriptor, temporary = tempfile.mkstemp(prefix='.workbook_', suffix='.xlsx',
                                             dir=output_dir)
    os.close(descriptor)
    try:
        workbook.save(temporary)
        os.replace(temporary, destination)
    finally:
        Path(temporary).unlink(missing_ok=True)
    return destination


def xml_cell(cell, strings):
    kind = cell.attrib.get('t')
    if cell.find(TAG + 'f') is not None:
        raise AssertionError('Unexpected Excel formula: ' + cell.attrib['r'])
    if kind == 'e':
        raise AssertionError('Unexpected Excel error cell: ' + cell.attrib['r'])
    value = cell.find(TAG + 'v')
    if kind == 'inlineStr':
        inline = cell.find(TAG + 'is')
        return '' if inline is None else ''.join(inline.itertext())
    if value is None or value.text is None:
        return None
    if kind == 's':
        return strings[int(value.text)]
    if kind == 'b':
        return value.text == '1'
    if kind == 'str':
        return value.text
    return float(value.text)


def assert_equal(actual, expected, context):
    if expected is None:
        equal = actual in (None, '')
    elif isinstance(expected, bool):
        equal = isinstance(actual, bool) and actual == expected
    elif isinstance(expected, (int, float)):
        equal = (isinstance(actual, (int, float)) and
                 not isinstance(actual, bool) and
                 math.isclose(actual, expected, rel_tol=2e-14, abs_tol=2e-14))
    else:
        equal = isinstance(actual, str) and actual == expected
    if not equal:
        raise AssertionError((context, expected, actual))


def verify_book(path, book):
    """Stream the ZIP/XML directly and compare every typed source data cell."""
    counts = dict(numeric_cells_checked=0, text_cells_checked=0,
                  blank_cells_checked=0)
    with ZipFile(path) as archive:
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            root = ET.fromstring(archive.read('xl/sharedStrings.xml'))
            strings = [''.join(node.itertext()) for node in root.findall(TAG + 'si')]
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        relationships = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {node.attrib['Id']: node.attrib['Target']
                   for node in relationships.findall('{' + PACKAGE_REL + '}Relationship')}
        sheets = {}
        for node in workbook.find(TAG + 'sheets'):
            target = targets[node.attrib['{' + DOC_REL + '}id']]
            sheets[node.attrib['name']] = posixpath.normpath(
                target.lstrip('/') if target.startswith('/') else 'xl/' + target)
        assert list(sheets) == ['Readme'] + [source['name'] for source in book['sheets']]

        # Validate the metadata sheet as well as all numerical source sheets.
        readme = ET.fromstring(archive.read(sheets['Readme']))
        readme_cells = {node.attrib['r']: xml_cell(node, strings)
                        for node in readme.iter(TAG + 'c')}
        assert_equal(readme_cells.get('A2'), book['title'], 'Readme title')
        for row, note in enumerate(book['notes'], 5):
            for column, expected in zip(['A', 'B'], note):
                assert_equal(readme_cells.get(column + str(row)), expected,
                             ('Readme', column, row))

        for source in book['sheets']:
            checked_rows = 0
            pane_found = False
            table_found = False
            with archive.open(sheets[source['name']]) as stream:
                for _, node in ET.iterparse(stream, events=['end']):
                    if node.tag == TAG + 'pane':
                        assert float(node.attrib.get('ySplit', 0)) == 5
                        pane_found = True
                    elif node.tag == TAG + 'tablePart':
                        table_found = True
                    elif node.tag == TAG + 'row':
                        row_number = int(node.attrib['r'])
                        cells = {cell.attrib['r']: xml_cell(cell, strings)
                                 for cell in node.findall(TAG + 'c')}
                        if row_number == 5:
                            for column, expected in enumerate(source['headers'], 1):
                                assert_equal(cells.get(get_column_letter(column) + '5'),
                                             expected, (book['id'], source['name'], 'header', column))
                        elif row_number >= 6:
                            expected_row = source['rows'][row_number - 6]
                            for column, expected in enumerate(expected_row, 1):
                                address = get_column_letter(column) + str(row_number)
                                assert_equal(cells.get(address), expected,
                                             (book['id'], source['name'], address))
                                key = ('blank_cells_checked' if expected is None else
                                       'numeric_cells_checked' if isinstance(expected, (int, float))
                                       and not isinstance(expected, bool) else 'text_cells_checked')
                                counts[key] += 1
                            checked_rows += 1
                        node.clear()
            assert checked_rows == len(source['rows'])
            assert pane_found and table_found
        for name in archive.namelist():
            if name.startswith('xl/tables/') and name.endswith('.xml'):
                table = ET.fromstring(archive.read(name))
                assert table.find(TAG + 'autoFilter') is not None
                columns = table.find(TAG + 'tableColumns')
                names = [column.attrib['name'].casefold() for column in columns]
                assert len(names) == len(set(names))
    return counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=HERE / 'workbook_inputs.json')
    parser.add_argument('--output-dir', type=Path, default=HERE / 'rebuilt')
    parser.add_argument('--book', action='append', default=[],
                        help='Workbook ID, e.g. S8 or F6; repeat or comma-separate IDs. Default: all 14.')
    parser.add_argument('--verify', action='store_true',
                        help='Independently reconcile every data cell using ZIP/XML readers.')
    args = parser.parse_args()
    inputs = json.loads(args.input.read_text())
    selected = {value for item in args.book for value in item.split(',') if value}
    known = {book['id'] for book in inputs['workbooks']}
    if selected - known:
        parser.error('Unknown workbook IDs: ' + ', '.join(sorted(selected - known)))
    books = [book for book in inputs['workbooks']
             if not selected or book['id'] in selected]
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for book in books:
        destination = args.output_dir / book['file']
        if destination.exists():
            parser.error(f'Refusing to overwrite existing workbook: {destination}')
    reports = []
    for book in books:
        path = build_book(book, args.output_dir)
        report = dict(id=book['id'], file=path.name, bytes=path.stat().st_size,
                      sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                      rows=sum(len(source['rows']) for source in book['sheets']))
        if args.verify:
            report.update(verify_book(path, book), all_source_values_reconciled=True,
                          value_mismatches=0, formula_cells=0, excel_error_cells=0)
        reports.append(report)
        print(json.dumps(report), flush=True)
    summary = dict(exporter='openpyxl write-only', source_input=args.input.name,
                   input_sha256=hashlib.sha256(args.input.read_bytes()).hexdigest(),
                   frozen_workbooks_modified=False, workbooks=reports)
    (args.output_dir / 'Portable_Workbook_Build_Report.json').write_text(
        json.dumps(summary, indent=2) + '\n')


if __name__ == '__main__':
    main()
